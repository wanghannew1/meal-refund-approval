"""
Tests for Excel parser module.
"""

from pathlib import Path

import openpyxl
import pytest
from collections import Counter
from decimal import Decimal

from app.excel_parser import (
    BonusRow,
    BonusSheet,
    InsuranceFeeItem,
    InsuranceSheet,
    InsuranceUnitRow,
    MealRefundRow,
    MealRefundSheet,
    _find_unit_row,
    _insurance_normalize_text,
    _is_insurance_monthly_header,
    _resolve_insurance_columns,
    list_insurance_sheets,
    parse_bonus_excel,
    parse_insurance_excel,
    parse_meal_refund_excel,
    list_sheets,
    SheetInfo,
    _xls_saved_selected_index,
)


class TestRealFixture:
    """Test suite using the real template fixture file."""

    def test_parse_title(self, real_fixture_path: Path):
        """Test title extraction from row 0 col 0."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        assert sheet.title == "松辽水利委员会综合服务中心餐费明细（06.15）"

    def test_parse_row_count(self, real_fixture_path: Path):
        """Test that exactly 17 data rows are parsed."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        assert len(sheet.rows) == 17

    def test_parse_first_row(self, real_fixture_path: Path):
        """Test first data row values."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        row = sheet.rows[0]
        assert row.serial == "1"
        assert row.name == "韩东旭"
        assert row.id_card == "210319197203110478"
        assert row.amount == "1860"

    def test_parse_last_row(self, real_fixture_path: Path):
        """Test last data row values."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        row = sheet.rows[-1]
        assert row.serial == "17"
        assert row.name == "韩魁双"
        assert row.id_card == "220104197804274415"
        assert row.amount == "1860"

    def test_parse_id_card_with_x(self, real_fixture_path: Path):
        """Test that ID cards ending in X are preserved."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        # Row 2 (张世宇) has ID ending in X
        row = sheet.rows[1]
        assert row.name == "张世宇"
        assert row.id_card.endswith("X")
        assert row.id_card == "22010519830911161X"

        # Row 16 (刘洋) also has ID ending in X
        row16 = sheet.rows[15]
        assert row16.name == "刘洋"
        assert row16.id_card == "22032319880710511X"

    def test_parse_total_amount(self, real_fixture_path: Path):
        """Test total amount normalization from ￥：31620."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        assert sheet.total_amount == "31620"

    def test_parse_total_chinese(self, real_fixture_path: Path):
        """Test Chinese uppercase amount text."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        assert sheet.total_chinese == "（人民币大写）叁万壹仟陆佰贰拾元"

    def test_parse_maker(self, real_fixture_path: Path):
        """Test maker name extraction from 制表：孙世多."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        assert sheet.maker == "孙世多"

    def test_no_warnings_for_valid_sheet(self, real_fixture_path: Path):
        """Test that a valid sheet produces no warnings."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        assert sheet.warnings == []

    def test_serial_normalization(self, real_fixture_path: Path):
        """Test that float serials (1.0, 2.0) are normalized to strings."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        for i, row in enumerate(sheet.rows, start=1):
            assert row.serial == str(i)

    def test_amount_normalization(self, real_fixture_path: Path):
        """Test that float amounts (1860.0) are normalized to clean strings."""
        sheet = parse_meal_refund_excel(real_fixture_path)
        for row in sheet.rows:
            assert row.amount == "1860"
            assert "." not in row.amount


class TestEdgeCases:
    """Test suite for edge cases using synthetic fixtures."""

    def test_missing_header_raises(self, tmp_path: Path):
        """Test that missing header row raises ValueError."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Some title"
        ws["A2"] = "Not a header"
        filepath = tmp_path / "no_header.xlsx"
        wb.save(filepath)

        with pytest.raises(ValueError, match="未找到表头行"):
            parse_meal_refund_excel(filepath)

    def test_no_data_rows_raises(self, tmp_path: Path):
        """Test that no data rows raises ValueError."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "身份证号码"
        ws["D2"] = "金额"
        # No data rows, directly 总计
        ws["A3"] = "总计"
        ws["D3"] = "0"
        filepath = tmp_path / "no_data.xlsx"
        wb.save(filepath)

        with pytest.raises(ValueError, match="未找到数据行"):
            parse_meal_refund_excel(filepath)

    def test_total_mismatch_warning(self, tmp_path: Path):
        """Test that sum != total produces a warning."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test Title"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "身份证号码"
        ws["D2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "张三"
        ws["C3"] = "110101199001011234"
        ws["D3"] = 500
        ws["A4"] = 2
        ws["B4"] = "李四"
        ws["C4"] = "110101199002022345"
        ws["D4"] = 600
        ws["A5"] = "总计"
        ws["B5"] = "（人民币大写）壹仟壹佰元"
        ws["D5"] = "￥：1100"  # Correct sum
        filepath = tmp_path / "mismatch.xlsx"
        wb.save(filepath)

        sheet = parse_meal_refund_excel(filepath)
        assert sheet.warnings == []  # 500+600=1100, matches

        # Now create one with wrong total
        ws["D5"] = "￥：9999"  # Wrong total
        wb.save(filepath)

        sheet2 = parse_meal_refund_excel(filepath)
        assert len(sheet2.warnings) == 1
        assert "不一致" in sheet2.warnings[0]

    def test_numeric_id_card_warning(self, tmp_path: Path):
        """Test that numeric ID card cell produces a warning."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "身份证号码"
        ws["D2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "张三"
        # ID card as number (will lose precision for 18 digits)
        ws["C3"] = 110101199001011234
        ws["D3"] = 500
        ws["A4"] = "总计"
        ws["D4"] = "￥：500"
        filepath = tmp_path / "numeric_id.xlsx"
        wb.save(filepath)

        sheet = parse_meal_refund_excel(filepath)
        assert len(sheet.warnings) >= 1
        assert "身份证号码" in sheet.warnings[0]

    def test_unsupported_format_raises(self, tmp_path: Path):
        """Test that unsupported file format raises ValueError."""
        filepath = tmp_path / "test.csv"
        filepath.write_text("a,b,c")

        with pytest.raises(ValueError, match="不支持的文件格式"):
            parse_meal_refund_excel(filepath)

    def test_empty_sheet_raises(self, tmp_path: Path):
        """Test that empty sheet raises ValueError."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        filepath = tmp_path / "empty.xlsx"
        wb.save(filepath)

        with pytest.raises(ValueError, match="未找到表头行"):
            parse_meal_refund_excel(filepath)

    def test_xlsx_format_works(self, tmp_path: Path):
        """Test that .xlsx format is parsed correctly."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "XLSX Title"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "身份证号码"
        ws["D2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "王五"
        ws["C3"] = "110101199003031234"
        ws["D3"] = 1000
        ws["A4"] = "总计"
        ws["B4"] = "壹仟元"
        ws["D4"] = "￥：1000"
        ws["A6"] = "制表：测试"
        filepath = tmp_path / "test.xlsx"
        wb.save(filepath)

        sheet = parse_meal_refund_excel(filepath)
        assert sheet.title == "XLSX Title"
        assert len(sheet.rows) == 1
        assert sheet.rows[0].name == "王五"
        assert sheet.total_amount == "1000"
        assert sheet.maker == "测试"


def _build_bonus_workbook(
    tmp_path: Path,
    filename: str = "bonus.xlsx",
    title: str = "科研部发放吉林大学2024年科研业绩津贴明细表（08.28）",
    data_rows: list[tuple] | None = None,
    bonus_total: float = 96000.0,
    tax_total: float = 19504.0,
    net_total: float = 76496.0,
    maker_text: str = "制表人：张朦",
    skip_totals: bool = False,
    skip_maker: bool = False,
    skip_header: bool = False,
) -> Path:
    """Build a synthetic bonus workbook mirroring the real template layout."""
    import openpyxl

    if data_rows is None:
        data_rows = [
            (1.0, 824282.0, "张三", "110101199001011234", 80000.0, 15000.0, 65000.0),
            (2.0, 824283.0, "李四", "110101199002022345", 16000.0, 4504.0, 11496.0),
        ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = title

    if not skip_header:
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金（单位：元）"
        ws["F2"] = "应纳个人所得税"
        ws["G2"] = "实发金额"

    row_idx = 3
    for dr in data_rows:
        for col_idx, val in enumerate(dr, start=1):
            ws.cell(row=row_idx, column=col_idx, value=val)
        row_idx += 1

    if not skip_totals:
        ws.cell(row=row_idx, column=1, value="合计")
        ws.cell(row=row_idx, column=5, value=bonus_total)
        ws.cell(row=row_idx, column=6, value=tax_total)
        ws.cell(row=row_idx, column=7, value=net_total)
        totals_idx = row_idx
        row_idx += 1

    if not skip_maker and not skip_totals:
        ws.cell(row=totals_idx + 1, column=7, value=maker_text)

    filepath = tmp_path / filename
    wb.save(filepath)
    return filepath


class TestBonusFixture:
    """Test suite for bonus spreadsheet parsing with synthetic fixtures."""

    def test_parse_title(self, tmp_path: Path):
        """Test title extraction from row 0 col 0."""
        fp = _build_bonus_workbook(tmp_path)
        sheet = parse_bonus_excel(fp)
        assert sheet.title == "科研部发放吉林大学2024年科研业绩津贴明细表（08.28）"

    def test_parse_row_count(self, tmp_path: Path):
        """Test that data rows are parsed correctly."""
        fp = _build_bonus_workbook(tmp_path)
        sheet = parse_bonus_excel(fp)
        assert len(sheet.rows) == 2

    def test_parse_first_row(self, tmp_path: Path):
        """Test first data row values."""
        fp = _build_bonus_workbook(tmp_path)
        sheet = parse_bonus_excel(fp)
        row = sheet.rows[0]
        assert row.serial == "1"
        assert row.employee_id == "824282"
        assert row.name == "张三"
        assert row.id_card == "110101199001011234"
        assert row.bonus == "80000"
        assert row.tax == "15000"
        assert row.net_amount == "65000"

    def test_parse_last_row(self, tmp_path: Path):
        """Test last data row values."""
        fp = _build_bonus_workbook(tmp_path)
        sheet = parse_bonus_excel(fp)
        row = sheet.rows[-1]
        assert row.serial == "2"
        assert row.employee_id == "824283"
        assert row.name == "李四"

    def test_float_employee_id_normalized(self, tmp_path: Path):
        """Test that float employee IDs are normalized to int strings."""
        fp = _build_bonus_workbook(tmp_path)
        sheet = parse_bonus_excel(fp)
        assert sheet.rows[0].employee_id == "824282"
        assert "." not in sheet.rows[0].employee_id

    def test_text_employee_id_zero_padding_preserved(self, tmp_path: Path):
        """Test that text employee IDs preserve zero-padding."""
        fp = _build_bonus_workbook(
            tmp_path,
            data_rows=[
                (1.0, "001234", "张三", "110101199001011234", 80000.0, 15000.0, 65000.0),
            ],
            bonus_total=80000.0,
            tax_total=15000.0,
            net_total=65000.0,
        )
        sheet = parse_bonus_excel(fp)
        assert sheet.rows[0].employee_id == "001234"

    def test_id_card_with_x_preserved(self, tmp_path: Path):
        """Test that ID cards ending in X are preserved."""
        fp = _build_bonus_workbook(
            tmp_path,
            data_rows=[
                (1.0, 824282.0, "张三", "22010519830911161X", 80000.0, 15000.0, 65000.0),
            ],
            bonus_total=80000.0,
            tax_total=15000.0,
            net_total=65000.0,
        )
        sheet = parse_bonus_excel(fp)
        assert sheet.rows[0].id_card == "22010519830911161X"

    def test_totals_extraction(self, tmp_path: Path):
        """Test totals extraction from 合计 row."""
        fp = _build_bonus_workbook(tmp_path)
        sheet = parse_bonus_excel(fp)
        assert sheet.bonus_total == "96000"
        assert sheet.tax_total == "19504"
        assert sheet.net_total == "76496"

    def test_maker_extraction_from_col6(self, tmp_path: Path):
        """Test maker extraction from col6 (制表人：张朦)."""
        fp = _build_bonus_workbook(tmp_path)
        sheet = parse_bonus_excel(fp)
        assert sheet.maker == "张朦"

    def test_maker_with_short_form(self, tmp_path: Path):
        """Test maker extraction with '制表：' (no 人)."""
        fp = _build_bonus_workbook(tmp_path, maker_text="制表：王五")
        sheet = parse_bonus_excel(fp)
        assert sheet.maker == "王五"

    def test_blank_row_skipping(self, tmp_path: Path):
        """Test that blank rows between data and totals are skipped."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "个税"
        ws["G2"] = "实发"
        ws["A3"] = 1.0
        ws["B3"] = 824282.0
        ws["C3"] = "张三"
        ws["D3"] = "110101199001011234"
        ws["E3"] = 50000.0
        ws["F3"] = 10000.0
        ws["G3"] = 40000.0
        # Blank row 4
        ws["A5"] = "合计"
        ws["E5"] = 50000.0
        ws["F5"] = 10000.0
        ws["G5"] = 40000.0
        ws["G6"] = "制表人：测试"
        fp = tmp_path / "blank_rows.xlsx"
        wb.save(fp)

        sheet = parse_bonus_excel(fp)
        assert len(sheet.rows) == 1
        assert sheet.rows[0].name == "张三"

    def test_no_header_raises(self, tmp_path: Path):
        """Test that missing header row raises ValueError."""
        fp = _build_bonus_workbook(tmp_path, skip_header=True)
        with pytest.raises(ValueError, match="未找到表头行"):
            parse_bonus_excel(fp)

    def test_no_data_rows_raises(self, tmp_path: Path):
        """Test that no data rows raises ValueError."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "个税"
        ws["G2"] = "实发"
        ws["A3"] = "合计"
        ws["E3"] = 0
        ws["F3"] = 0
        ws["G3"] = 0
        fp = tmp_path / "no_data.xlsx"
        wb.save(fp)

        with pytest.raises(ValueError, match="未找到数据行"):
            parse_bonus_excel(fp)

    def test_totals_mismatch_warning(self, tmp_path: Path):
        """Test that sum != total produces a warning."""
        fp = _build_bonus_workbook(
            tmp_path,
            data_rows=[
                (1.0, 824282.0, "张三", "110101199001011234", 80000.0, 15000.0, 65000.0),
                (2.0, 824283.0, "李四", "110101199002022345", 16000.0, 4504.0, 11496.0),
            ],
            bonus_total=99999.0,
            tax_total=19504.0,
            net_total=76496.0,
        )
        sheet = parse_bonus_excel(fp)
        bonus_warnings = [w for w in sheet.warnings if "奖金明细合计" in w]
        assert len(bonus_warnings) == 1
        assert "不一致" in bonus_warnings[0]

    def test_per_row_bonus_tax_net_warning(self, tmp_path: Path):
        """Test per-row warning when bonus - tax != net_amount."""
        fp = _build_bonus_workbook(
            tmp_path,
            data_rows=[
                (1.0, 824282.0, "张三", "110101199001011234", 80000.0, 15000.0, 99999.0),
            ],
            bonus_total=80000.0,
            tax_total=15000.0,
            net_total=99999.0,
        )
        sheet = parse_bonus_excel(fp)
        row_warnings = [w for w in sheet.warnings if "实发金额 与 奖金-个人所得税 不一致" in w]
        assert len(row_warnings) == 1
        assert "第1行" in row_warnings[0]

    def test_unsupported_format_raises(self, tmp_path: Path):
        """Test that unsupported file format raises ValueError."""
        filepath = tmp_path / "test.csv"
        filepath.write_text("a,b,c")
        with pytest.raises(ValueError, match="不支持的文件格式"):
            parse_bonus_excel(filepath)


class TestBonusAdjustmentColumn:
    """Test suite for the optional 个人其他调整 (adjustment) column."""

    def test_adjustment_column_absent(self, tmp_path: Path):
        """When adjustment column is absent, all rows have adjustment='0'."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "应纳个人所得税"
        ws["G2"] = "实发金额"
        ws["A3"] = 1.0
        ws["B3"] = 824282.0
        ws["C3"] = "张三"
        ws["D3"] = "110101199001011234"
        ws["E3"] = 50000.0
        ws["F3"] = 10000.0
        ws["G3"] = 40000.0
        ws["A4"] = "合计"
        ws["E4"] = 50000.0
        ws["F4"] = 10000.0
        ws["G4"] = 40000.0
        ws["G5"] = "制表人：测试"
        fp = tmp_path / "no_adj.xlsx"
        wb.save(fp)

        sheet = parse_bonus_excel(fp)
        assert len(sheet.rows) == 1
        assert sheet.rows[0].adjustment == "0"
        assert sheet.adjustment_total == "0"
        adj_warnings = [w for w in sheet.warnings if "个人其他调整" in w]
        assert len(adj_warnings) == 0

    def test_adjustment_column_at_end(self, tmp_path: Path):
        """When adjustment column is at the end, values are parsed."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "应纳个人所得税"
        ws["G2"] = "实发金额"
        ws["H2"] = "个人其他调整"
        ws["A3"] = 1.0
        ws["B3"] = 824282.0
        ws["C3"] = "张三"
        ws["D3"] = "110101199001011234"
        ws["E3"] = 50000.0
        ws["F3"] = 10000.0
        ws["G3"] = 40000.0
        ws["H3"] = 500.0
        ws["A4"] = "合计"
        ws["E4"] = 50000.0
        ws["F4"] = 10000.0
        ws["G4"] = 40000.0
        ws["H4"] = 500.0
        ws["H5"] = "制表人：测试"
        fp = tmp_path / "adj_end.xlsx"
        wb.save(fp)

        sheet = parse_bonus_excel(fp)
        assert sheet.rows[0].adjustment == "500"
        assert sheet.adjustment_total == "500"

    def test_adjustment_column_between_tax_and_net(self, tmp_path: Path):
        """When adjustment is between tax and net, all fields parse correctly."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "应纳个人所得税"
        ws["G2"] = "个人其他调整"
        ws["H2"] = "实发金额"
        ws["A3"] = 1.0
        ws["B3"] = 824282.0
        ws["C3"] = "张三"
        ws["D3"] = "110101199001011234"
        ws["E3"] = 50000.0
        ws["F3"] = 10000.0
        ws["G3"] = 500.0
        ws["H3"] = 39500.0
        ws["A4"] = "合计"
        ws["E4"] = 50000.0
        ws["F4"] = 10000.0
        ws["G4"] = 500.0
        ws["H4"] = 39500.0
        ws["H5"] = "制表人：测试"
        fp = tmp_path / "adj_mid.xlsx"
        wb.save(fp)

        sheet = parse_bonus_excel(fp)
        row = sheet.rows[0]
        assert row.bonus == "50000"
        assert row.tax == "10000"
        assert row.adjustment == "500"
        assert row.net_amount == "39500"
        assert sheet.adjustment_total == "500"

    def test_adjustment_blank_cell_becomes_zero(self, tmp_path: Path):
        """Blank adjustment cell becomes '0'."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "应纳个人所得税"
        ws["G2"] = "个人其他调整"
        ws["H2"] = "实发金额"
        ws["A3"] = 1.0
        ws["B3"] = 824282.0
        ws["C3"] = "张三"
        ws["D3"] = "110101199001011234"
        ws["E3"] = 50000.0
        ws["F3"] = 10000.0
        ws["G3"] = None
        ws["H3"] = 40000.0
        ws["A4"] = "合计"
        ws["E4"] = 50000.0
        ws["F4"] = 10000.0
        ws["G4"] = 0.0
        ws["H4"] = 40000.0
        ws["H5"] = "制表人：测试"
        fp = tmp_path / "adj_blank.xlsx"
        wb.save(fp)

        sheet = parse_bonus_excel(fp)
        assert sheet.rows[0].adjustment == "0"

    def test_blank_money_and_input_cells_do_not_leak_none(self, tmp_path: Path):
        """Blank money cells become '0'; blank input cells become '' (never 'None')."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "应纳个人所得税"
        ws["G2"] = "实发金额"
        ws["C3"] = "张三"
        ws["A4"] = "合计"
        ws["H5"] = "制表人：测试"
        fp = tmp_path / "blank_cells.xlsx"
        wb.save(fp)

        row = parse_bonus_excel(fp).rows[0]
        assert row.bonus == "0"
        assert row.tax == "0"
        assert row.net_amount == "0"
        assert row.adjustment == "0"
        assert row.serial == ""
        assert row.employee_id == ""
        assert row.id_card == ""
        assert "None" not in (row.bonus, row.tax, row.net_amount, row.serial, row.employee_id, row.id_card)

    def test_adjustment_negative_value(self, tmp_path: Path):
        """Negative adjustment value is preserved."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "应纳个人所得税"
        ws["G2"] = "个人其他调整"
        ws["H2"] = "实发金额"
        ws["A3"] = 1.0
        ws["B3"] = 824282.0
        ws["C3"] = "张三"
        ws["D3"] = "110101199001011234"
        ws["E3"] = 50000.0
        ws["F3"] = 10000.0
        ws["G3"] = -500.0
        ws["H3"] = 40500.0
        ws["A4"] = "合计"
        ws["E4"] = 50000.0
        ws["F4"] = 10000.0
        ws["G4"] = -500.0
        ws["H4"] = 40500.0
        ws["H5"] = "制表人：测试"
        fp = tmp_path / "adj_neg.xlsx"
        wb.save(fp)

        sheet = parse_bonus_excel(fp)
        assert sheet.rows[0].adjustment == "-500"

    def test_adjustment_total_mismatch_warning(self, tmp_path: Path):
        """Adjustment sum mismatch produces a warning."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "应纳个人所得税"
        ws["G2"] = "个人其他调整"
        ws["H2"] = "实发金额"
        ws["A3"] = 1.0
        ws["B3"] = 824282.0
        ws["C3"] = "张三"
        ws["D3"] = "110101199001011234"
        ws["E3"] = 50000.0
        ws["F3"] = 10000.0
        ws["G3"] = 500.0
        ws["H3"] = 39500.0
        ws["A4"] = "合计"
        ws["E4"] = 50000.0
        ws["F4"] = 10000.0
        ws["G4"] = 9999.0
        ws["H4"] = 39500.0
        ws["H5"] = "制表人：测试"
        fp = tmp_path / "adj_mismatch.xlsx"
        wb.save(fp)

        sheet = parse_bonus_excel(fp)
        adj_warnings = [w for w in sheet.warnings if "个人其他调整明细合计" in w]
        assert len(adj_warnings) == 1
        assert "不一致" in adj_warnings[0]

    def test_per_row_net_check_skipped_with_nonzero_adjustment(self, tmp_path: Path):
        """Per-row net check is skipped when adjustment is non-zero."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "应纳个人所得税"
        ws["G2"] = "个人其他调整"
        ws["H2"] = "实发金额"
        ws["A3"] = 1.0
        ws["B3"] = 824282.0
        ws["C3"] = "张三"
        ws["D3"] = "110101199001011234"
        ws["E3"] = 50000.0
        ws["F3"] = 10000.0
        ws["G3"] = 500.0
        ws["H3"] = 99999.0
        ws["A4"] = "合计"
        ws["E4"] = 50000.0
        ws["F4"] = 10000.0
        ws["G4"] = 500.0
        ws["H4"] = 99999.0
        ws["H5"] = "制表人：测试"
        fp = tmp_path / "adj_skip_check.xlsx"
        wb.save(fp)

        sheet = parse_bonus_excel(fp)
        net_warnings = [w for w in sheet.warnings if "实发金额 与 奖金-个人所得税 不一致" in w]
        assert len(net_warnings) == 0

    def test_per_row_net_check_active_with_zero_adjustment(self, tmp_path: Path):
        """Per-row net check is active when adjustment is zero."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Title"
        ws["A2"] = "序号"
        ws["B2"] = "工号"
        ws["C2"] = "姓名"
        ws["D2"] = "身份证号"
        ws["E2"] = "奖金"
        ws["F2"] = "应纳个人所得税"
        ws["G2"] = "个人其他调整"
        ws["H2"] = "实发金额"
        ws["A3"] = 1.0
        ws["B3"] = 824282.0
        ws["C3"] = "张三"
        ws["D3"] = "110101199001011234"
        ws["E3"] = 50000.0
        ws["F3"] = 10000.0
        ws["G3"] = 0.0
        ws["H3"] = 99999.0
        ws["A4"] = "合计"
        ws["E4"] = 50000.0
        ws["F4"] = 10000.0
        ws["G4"] = 0.0
        ws["H4"] = 99999.0
        ws["H5"] = "制表人：测试"
        fp = tmp_path / "adj_zero_check.xlsx"
        wb.save(fp)

        sheet = parse_bonus_excel(fp)
        net_warnings = [w for w in sheet.warnings if "实发金额 与 奖金-个人所得税 不一致" in w]
        assert len(net_warnings) == 1
        assert "第1行" in net_warnings[0]


def _write_union_fee_workbook(
    path: Path,
    total_amount: int = 45,
    include_total_row: bool = True,
    include_signature_row: bool = True,
) -> None:
    """Write a 5-column 工会会费 sheet shaped like the real template.

    Layout mirrors the real file (title, 单位行, header, 3 data rows,
    合计 row, 说明 row, blank row, signature row). The 合计 column at
    index 4 duplicates the 金额 value and must be ignored by the parser.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "XX单位2026年06月派遣员工工会会费转入彩虹公司工会账户明细（06.03）"
    ws["E2"] = "单位：元"
    ws["A3"] = "序号"
    ws["B3"] = "姓名"
    ws["C3"] = "身份证号"
    ws["D3"] = "金额"
    ws["E3"] = "合计"
    ws["A4"] = 1
    ws["B4"] = "陈一"
    ws["C4"] = "110101197001011234"
    ws["D4"] = 15
    ws["E4"] = 15
    ws["A5"] = 2
    ws["B5"] = "陈二"
    ws["C5"] = "11010119710202234X"
    ws["D5"] = 15
    ws["E5"] = 15
    ws["A6"] = 3
    ws["B6"] = "陈三"
    ws["C6"] = "110101197203033456"
    ws["D6"] = 15
    ws["E6"] = 15
    if include_total_row:
        ws["A7"] = "合计"
        ws["D7"] = total_amount
        ws["E7"] = total_amount
    ws["A8"] = "说明：现申请将该单位工会会费转入彩虹公司工会账户。"
    # row 9 intentionally blank
    if include_signature_row:
        ws["A10"] = "总经理签字："
        ws["B10"] = "部长签字："
        ws["D10"] = "财务审核："
        ws["E10"] = "制表：郑文琦"
    wb.save(path)


class TestUnionFeeSheet:
    """Test suite for the 5-column 工会会费 (union fee) spreadsheet format."""

    def test_union_fee_total_row_detected(self, tmp_path: Path):
        """Total row detected via col0=='合计'; maker from col index 4."""
        fp = tmp_path / "union_fee.xlsx"
        _write_union_fee_workbook(fp)
        sheet = parse_meal_refund_excel(fp)
        assert sheet.total_amount == "45"
        assert sheet.total_chinese == ""
        assert sheet.maker == "郑文琦"

    def test_union_fee_rows_and_extra_column_ignored(self, tmp_path: Path):
        """Exactly 3 data rows; amount from col3, not the 合计 col4."""
        fp = tmp_path / "union_fee.xlsx"
        _write_union_fee_workbook(fp)
        sheet = parse_meal_refund_excel(fp)
        assert len(sheet.rows) == 3
        row0 = sheet.rows[0]
        assert row0 == MealRefundRow(
            serial="1",
            name="陈一",
            id_card="110101197001011234",
            amount="15",
        )
        assert sheet.rows[2].amount == "15"

    def test_union_fee_no_bogus_signature_rows(self, tmp_path: Path):
        """Signature/说明 rows must not leak into the parsed data."""
        fp = tmp_path / "union_fee.xlsx"
        _write_union_fee_workbook(fp)
        sheet = parse_meal_refund_excel(fp)
        assert len(sheet.rows) == 3
        assert not any("签字" in r.serial for r in sheet.rows)
        assert not any(r.name == "部长签字：" for r in sheet.rows)

    def test_union_fee_sum_mismatch_warns(self, tmp_path: Path):
        """合计 total != sum(rows) produces exactly one '不一致' warning."""
        fp = tmp_path / "union_fee_mismatch.xlsx"
        _write_union_fee_workbook(fp, total_amount=99)
        sheet = parse_meal_refund_excel(fp)
        mismatch_warnings = [w for w in sheet.warnings if "不一致" in w]
        assert len(mismatch_warnings) == 1

    def test_meal_total_row_still_works(self, tmp_path: Path):
        """Regression guard: 4-column 返餐费 format with col0=='总计' still parses."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test Meal"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "身份证号码"
        ws["D2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "张三"
        ws["C3"] = "110101199001011234"
        ws["D3"] = 60
        ws["A4"] = 2
        ws["B4"] = "李四"
        ws["C4"] = "110101199002022345"
        ws["D4"] = 60
        ws["A5"] = "总计"
        ws["B5"] = "（人民币大写）壹佰贰拾元"
        ws["D5"] = "￥：120"
        ws["A7"] = "制表：测试"
        fp = tmp_path / "meal.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp)
        assert sheet.total_amount == "120"
        assert sheet.total_chinese == "（人民币大写）壹佰贰拾元"
        assert len(sheet.rows) == 2

    def test_signature_guard_stops_without_total_row(self, tmp_path: Path):
        """Signature row stops parsing even when no total row exists."""
        fp = tmp_path / "no_total.xlsx"
        _write_union_fee_workbook(fp, include_total_row=False)
        sheet = parse_meal_refund_excel(fp)
        assert len(sheet.rows) == 3
        assert sheet.total_amount == ""
        assert not any("签字" in r.serial for r in sheet.rows)


class TestThreeColumnNoID:
    """Test suite for 3-column layout without ID card column."""

    def test_three_col_no_id_parses(self, tmp_path: Path):
        """3-col layout (序号/姓名/金额) parses correctly with id_card=''."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test Title"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "张三"
        ws["C3"] = 100.50
        ws["A4"] = 2
        ws["B4"] = "李四"
        ws["C4"] = 200
        ws["A5"] = "总计"
        ws["B5"] = "叁佰元伍角"
        ws["C5"] = 300.50
        ws["A7"] = "制表：测试"
        fp = tmp_path / "three_col.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp)
        assert len(sheet.rows) == 2
        assert sheet.rows[0].serial == "1"
        assert sheet.rows[0].name == "张三"
        assert sheet.rows[0].id_card == ""
        assert sheet.rows[0].amount == "100.5"
        assert sheet.rows[1].name == "李四"
        assert sheet.rows[1].id_card == ""
        assert sheet.rows[1].amount == "200"
        assert sheet.total_amount == "300.5"
        assert sheet.maker == "测试"

    def test_three_col_no_id_card_warnings(self, tmp_path: Path):
        """3-col layout produces no ID-card-related warnings."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "张三"
        ws["C3"] = 100
        ws["A4"] = "总计"
        ws["C4"] = 100
        fp = tmp_path / "no_id.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp)
        assert not any("身份证号码为数字格式" in w for w in sheet.warnings)


class TestMultiSheetWorkbooks:
    """Test suite for multi-sheet workbook support."""

    def test_list_sheets_xlsx(self, tmp_path: Path):
        """list_sheets returns sheet names and header validity for .xlsx."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1["A1"] = "Title"
        ws1["A2"] = "序号"
        ws1["B2"] = "姓名"
        ws1["C2"] = "金额"
        ws1["A3"] = 1
        ws1["B3"] = "张三"
        ws1["C3"] = 100

        ws2 = wb.create_sheet("Sheet2")
        ws2["A1"] = "No header here"

        fp = tmp_path / "multi.xlsx"
        wb.save(fp)

        sheets = list_sheets(fp)
        assert len(sheets) == 2
        assert sheets[0].name == "Sheet1"
        assert sheets[0].has_valid_header is True
        assert sheets[1].name == "Sheet2"
        assert sheets[1].has_valid_header is False
        assert sheets[0].is_default is True
        assert sheets[1].is_default is False

    def test_list_sheets_xls(self, tmp_path: Path):
        """list_sheets works for .xls files."""
        pytest.importorskip("xlwt")
        import xlwt

        wb = xlwt.Workbook()
        ws1 = wb.add_sheet("Sheet1")
        ws1.write(0, 0, "序号")
        ws1.write(0, 1, "姓名")
        ws1.write(0, 2, "金额")
        ws1.write(1, 0, 1)
        ws1.write(1, 1, "张三")
        ws1.write(1, 2, 100)

        ws2 = wb.add_sheet("Sheet2")
        ws2.write(0, 0, "No header")

        fp = tmp_path / "multi.xls"
        wb.save(fp)

        sheets = list_sheets(fp)
        assert len(sheets) == 2
        assert sheets[0].name == "Sheet1"
        assert sheets[0].has_valid_header is True
        assert sheets[1].name == "Sheet2"
        assert sheets[1].has_valid_header is False
        assert sheets[0].is_default is False
        assert sheets[1].is_default is True

    def test_xlsx_default_active_sheet(self, tmp_path: Path):
        """For .xlsx, default sheet is the active sheet."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "InactiveSheet"
        ws1["A1"] = "No header"

        ws2 = wb.create_sheet("ActiveSheet")
        ws2["A1"] = "Title"
        ws2["A2"] = "序号"
        ws2["B2"] = "姓名"
        ws2["C2"] = "金额"
        ws2["A3"] = 1
        ws2["B3"] = "张三"
        ws2["C3"] = 100
        ws2["A4"] = "总计"
        ws2["C4"] = 100
        wb.active = ws2

        fp = tmp_path / "active.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp)
        assert sheet.rows[0].name == "张三"

    def test_xls_default_last_sheet(self, tmp_path: Path):
        """For .xls, default sheet is the last sheet."""
        pytest.importorskip("xlwt")
        import xlwt

        wb = xlwt.Workbook()
        ws1 = wb.add_sheet("FirstSheet")
        ws1.write(0, 0, "No header")

        ws2 = wb.add_sheet("LastSheet")
        ws2.write(0, 0, "序号")
        ws2.write(0, 1, "姓名")
        ws2.write(0, 2, "金额")
        ws2.write(1, 0, 1)
        ws2.write(1, 1, "张三")
        ws2.write(1, 2, 100)
        ws2.write(2, 0, "总计")
        ws2.write(2, 2, 100)

        fp = tmp_path / "last.xls"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp)
        assert sheet.rows[0].name == "张三"

    def test_sheet_selection_by_name(self, tmp_path: Path):
        """Parse a specific sheet by name."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1["A1"] = "Title1"
        ws1["A2"] = "序号"
        ws1["B2"] = "姓名"
        ws1["C2"] = "金额"
        ws1["A3"] = 1
        ws1["B3"] = "张三"
        ws1["C3"] = 100
        ws1["A4"] = "总计"
        ws1["C4"] = 100

        ws2 = wb.create_sheet("Sheet2")
        ws2["A1"] = "Title2"
        ws2["A2"] = "序号"
        ws2["B2"] = "姓名"
        ws2["C2"] = "金额"
        ws2["A3"] = 1
        ws2["B3"] = "李四"
        ws2["C3"] = 200
        ws2["A4"] = "总计"
        ws2["C4"] = 200

        fp = tmp_path / "select.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp, sheet_name="Sheet2")
        assert sheet.title == "Title2"
        assert sheet.rows[0].name == "李四"
        assert sheet.rows[0].amount == "200"

    def test_sheet_selection_by_index(self, tmp_path: Path):
        """Parse a specific sheet by index."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1["A1"] = "Title1"
        ws1["A2"] = "序号"
        ws1["B2"] = "姓名"
        ws1["C2"] = "金额"
        ws1["A3"] = 1
        ws1["B3"] = "张三"
        ws1["C3"] = 100
        ws1["A4"] = "总计"
        ws1["C4"] = 100

        ws2 = wb.create_sheet("Sheet2")
        ws2["A1"] = "Title2"
        ws2["A2"] = "序号"
        ws2["B2"] = "姓名"
        ws2["C2"] = "金额"
        ws2["A3"] = 1
        ws2["B3"] = "李四"
        ws2["C3"] = 200
        ws2["A4"] = "总计"
        ws2["C4"] = 200

        fp = tmp_path / "index.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp, sheet_index=1)
        assert sheet.title == "Title2"
        assert sheet.rows[0].name == "李四"


class TestDefaultSheetSelection:
    """SheetInfo.is_default must match the parser's no-argument default sheet.

    Agreed rule: .xlsx -> the workbook's ACTIVE sheet (unconditional, even if
    it lacks a valid header); .xls -> the LAST sheet.
    """

    @staticmethod
    def _write_active_default_workbook(tmp_path: Path) -> Path:
        """Two valid-header sheets; Sheet2 saved as the active tab (Sheet1 is the decoy)."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1["A1"] = "Title1"
        ws1["A2"] = "序号"
        ws1["B2"] = "姓名"
        ws1["C2"] = "金额"
        ws1["A3"] = 1
        ws1["B3"] = "不该选中"
        ws1["C3"] = 100
        ws1["A4"] = "总计"
        ws1["C4"] = 100

        ws2 = wb.create_sheet("Sheet2")
        ws2["A1"] = "Title2"
        ws2["A2"] = "序号"
        ws2["B2"] = "姓名"
        ws2["C2"] = "金额"
        ws2["A3"] = 1
        ws2["B3"] = "激活表数据"
        ws2["C3"] = 200
        ws2["A4"] = "总计"
        ws2["C4"] = 200

        wb.active = 1
        fp = tmp_path / "active_default.xlsx"
        wb.save(fp)
        return fp

    def test_xlsx_flags_active_sheet_and_parses_it_by_default(self, tmp_path: Path):
        """list_sheets flags the active sheet; no-arg parse reads it, not sheet 0."""
        fp = self._write_active_default_workbook(tmp_path)

        infos = list_sheets(fp)
        assert [s.name for s in infos if s.is_default] == ["Sheet2"]

        sheet = parse_meal_refund_excel(fp)
        assert sheet.rows[0].name == "激活表数据"

        decoy = parse_meal_refund_excel(fp, sheet_name="Sheet1")
        assert decoy.rows[0].name == "不该选中"

    def test_default_flag_matches_parser_no_arg_rows(self, tmp_path: Path):
        """Parsing by the flagged name yields the same rows as parsing with sheet_name=None."""
        fp = self._write_active_default_workbook(tmp_path)

        infos = list_sheets(fp)
        default_name = next(s.name for s in infos if s.is_default)

        by_name = parse_meal_refund_excel(fp, sheet_name=default_name)
        by_default = parse_meal_refund_excel(fp)
        assert by_name.rows == by_default.rows

    def test_xls_flags_last_sheet(self, real_fixture_path: Path):
        """The .xls last-sheet rule is flagged (single-sheet fixture -> its only sheet)."""
        infos = list_sheets(real_fixture_path)
        assert infos
        assert [s.name for s in infos if s.is_default] == [infos[-1].name]


class TestXlsSavedSelectedIndex:
    """Test suite for _xls_saved_selected_index helper (WINDOW1 record reader)."""

    def test_synthetic_window1_record(self, tmp_path: Path):
        """Helper extracts itabCur from a synthetic WINDOW1 record."""
        import struct

        # Build a minimal byte blob with a valid WINDOW1 record:
        # type 0x003D (little-endian: b'\x3d\x00'), length 0x0012 (b'\x12\x00'),
        # then 18 bytes of data with itabCur at offset 10 set to 1.
        header = b"\x3d\x00\x12\x00"
        # 9 uint16 fields: xWn, yWn, dxWn, dyWn, grbit, itabCur, itabFirst, ctabSel, wTabRatio
        # Set itabCur (field index 5) to 1, others to 0.
        data = struct.pack("<HHHHHHHHH", 0, 0, 0, 0, 0, 1, 0, 0, 0)
        blob = header + data
        fp = tmp_path / "synthetic.xls"
        fp.write_bytes(blob)

        result = _xls_saved_selected_index(fp)
        assert result == 1

    def test_garbage_file_returns_none(self, tmp_path: Path):
        """Helper returns None for a garbage file."""
        fp = tmp_path / "garbage.xls"
        fp.write_bytes(b"not an xls")

        result = _xls_saved_selected_index(fp)
        assert result is None

    def test_fallback_when_helper_returns_none(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
        """When helper returns None, is_default lands on the last sheet."""
        pytest.importorskip("xlwt")
        import xlwt

        wb_xls = xlwt.Workbook()
        ws1_xls = wb_xls.add_sheet("Sheet1")
        ws1_xls.write(0, 0, "序号")
        ws1_xls.write(0, 1, "姓名")
        ws1_xls.write(0, 2, "金额")
        ws1_xls.write(1, 0, 1)
        ws1_xls.write(1, 1, "张三")
        ws1_xls.write(1, 2, 100)

        ws2_xls = wb_xls.add_sheet("Sheet2")
        ws2_xls.write(0, 0, "No header")

        fp_xls = tmp_path / "fallback.xls"
        wb_xls.save(fp_xls)

        monkeypatch.setattr(
            "app.excel_parser._xls_saved_selected_index",
            lambda path: None
        )

        infos = list_sheets(fp_xls)
        assert [s.name for s in infos if s.is_default] == ["Sheet2"]

    def test_consistency_flagged_default_matches_parser(self, tmp_path: Path):
        """For any multi-sheet fixture, the name flagged is_default by list_sheets
        equals what a no-arg parse reads (extends TestDefaultSheetSelection pattern)."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1["A1"] = "Title1"
        ws1["A2"] = "序号"
        ws1["B2"] = "姓名"
        ws1["C2"] = "金额"
        ws1["A3"] = 1
        ws1["B3"] = "张三"
        ws1["C3"] = 100
        ws1["A4"] = "总计"
        ws1["C4"] = 100

        ws2 = wb.create_sheet("Sheet2")
        ws2["A1"] = "Title2"
        ws2["A2"] = "序号"
        ws2["B2"] = "姓名"
        ws2["C2"] = "金额"
        ws2["A3"] = 1
        ws2["B3"] = "李四"
        ws2["C3"] = 200
        ws2["A4"] = "总计"
        ws2["C4"] = 200

        wb.active = 1  # Sheet2 is active
        fp = tmp_path / "consistency.xlsx"
        wb.save(fp)

        infos = list_sheets(fp)
        default_name = next(s.name for s in infos if s.is_default)
        by_default = parse_meal_refund_excel(fp)
        by_name = parse_meal_refund_excel(fp, sheet_name=default_name)
        assert by_name.rows == by_default.rows


class TestDuplicateNames:
    """Test suite for duplicate name detection."""

    def test_duplicate_names_kept_independent(self, tmp_path: Path):
        """Duplicate names are kept as separate rows, never merged."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "张三"
        ws["C3"] = 100
        ws["A4"] = 2
        ws["B4"] = "张三"
        ws["C4"] = 200
        ws["A5"] = 3
        ws["B5"] = "李四"
        ws["C5"] = 300
        ws["A6"] = "总计"
        ws["C6"] = 600
        fp = tmp_path / "dup.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp)
        assert len(sheet.rows) == 3
        assert sheet.rows[0].name == "张三"
        assert sheet.rows[0].amount == "100"
        assert sheet.rows[1].name == "张三"
        assert sheet.rows[1].amount == "200"
        assert sheet.rows[2].name == "李四"

    def test_duplicate_names_warning(self, tmp_path: Path):
        """Duplicate names produce a warning listing names and counts."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "张三"
        ws["C3"] = 100
        ws["A4"] = 2
        ws["B4"] = "张三"
        ws["C4"] = 200
        ws["A5"] = 3
        ws["B5"] = "李四"
        ws["C5"] = 300
        ws["A6"] = 4
        ws["B6"] = "李四"
        ws["C6"] = 400
        ws["A7"] = "总计"
        ws["C7"] = 1000
        fp = tmp_path / "dup_warn.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp)
        dup_warnings = [w for w in sheet.warnings if "重复姓名" in w]
        assert len(dup_warnings) == 1
        assert "张三(2次)" in dup_warnings[0]
        assert "李四(2次)" in dup_warnings[0]
        assert "各行独立保留" in dup_warnings[0]


class TestSkippedRows:
    """Test suite for skipped-row accounting."""

    def test_skipped_blank_rows_counted(self, tmp_path: Path):
        """Blank rows are counted and reported in warnings."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "张三"
        ws["C3"] = 100
        # Blank row 4
        ws["A5"] = 2
        ws["B5"] = "李四"
        ws["C5"] = 200
        ws["A6"] = "总计"
        ws["C6"] = 300
        fp = tmp_path / "blank.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp)
        assert len(sheet.rows) == 2
        skip_warnings = [w for w in sheet.warnings if "跳过" in w]
        assert len(skip_warnings) == 1
        assert "1 行" in skip_warnings[0]

    def test_skipped_nameless_rows_counted(self, tmp_path: Path):
        """Rows with no name are counted and reported."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "张三"
        ws["C3"] = 100
        ws["A4"] = 2
        ws["B4"] = ""
        ws["C4"] = 200
        ws["A5"] = 3
        ws["B5"] = "李四"
        ws["C5"] = 300
        ws["A6"] = "总计"
        ws["C6"] = 400
        fp = tmp_path / "nameless.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp)
        assert len(sheet.rows) == 2
        skip_warnings = [w for w in sheet.warnings if "跳过" in w]
        assert len(skip_warnings) == 1
        assert "1 行" in skip_warnings[0]


class TestEmptyIDCard:
    """Test suite for empty ID card handling."""

    def test_empty_id_card_cell_becomes_empty_string(self, tmp_path: Path):
        """Empty/None ID card cell becomes '' (never 'None')."""
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test"
        ws["A2"] = "序号"
        ws["B2"] = "姓名"
        ws["C2"] = "身份证号码"
        ws["D2"] = "金额"
        ws["A3"] = 1
        ws["B3"] = "张三"
        ws["C3"] = None
        ws["D3"] = 100
        ws["A4"] = "总计"
        ws["D4"] = 100
        fp = tmp_path / "empty_id.xlsx"
        wb.save(fp)

        sheet = parse_meal_refund_excel(fp)
        assert sheet.rows[0].id_card == ""
        assert sheet.rows[0].id_card != "None"
        assert "None" not in sheet.rows[0].id_card


class TestRealFixtureThreeCol:
    """Test suite using the real 3-col fixture file."""

    def test_real_three_col_fixture(self):
        """Test parsing the real 3-col fixture (博众返餐费.xlsx), pinned to Sheet1.

        Sheet1 is pinned because the workbook's active tab is user-controlled;
        the no-arg default (active sheet) is covered by TestDefaultSheetSelection.
        """
        from pathlib import Path
        from decimal import Decimal

        fixture_path = Path("/home/ubuntu/excel_example/mealrefund/博众返餐费.xlsx")
        if not fixture_path.exists():
            pytest.skip("3-col fixture not available")

        sheet = parse_meal_refund_excel(fixture_path, sheet_name="Sheet1")
        assert len(sheet.rows) == 515
        assert sheet.rows[0].name == "李宏奇"
        assert Decimal(sheet.rows[0].amount) == Decimal("0.04")
        assert sheet.rows[0].id_card == ""
        assert sheet.rows[-1].name == "李妍宁"
        assert Decimal(sheet.rows[-1].amount) == Decimal("0.33")
        assert Decimal(sheet.total_amount) == Decimal("16153.95")
        assert sheet.maker == "张三"
        assert not any("身份证号码为数字格式" in w for w in sheet.warnings)
        assert not any("不一致" in w for w in sheet.warnings)


class _InsuranceWorkbook:
    @staticmethod
    def make(tmp_path: Path, rows: list[list[object]], sheet_name: str = "Sheet1") -> Path:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name
        for r, row in enumerate(rows, start=1):
            for c, val in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=val)
        fp = tmp_path / f"{sheet_name}.xlsx"
        wb.save(fp)
        return fp


class TestInsuranceHeaderDetection:
    def _final_names(self, resolved, rows):
        unit_row_idx = _find_unit_row(rows)
        name_row_idx = unit_row_idx + 1
        normalized = [
            _insurance_normalize_text(rows[name_row_idx][col]) for _, col in resolved[2]
        ]
        counts = Counter(normalized)
        return [
            f"{group}-{name}" if counts[name] >= 2 and group else name
            for (group, col), name in zip(resolved[2], normalized)
        ]

    def test_medical_base_header_true(self, tmp_path):
        rows = [
            ["2026年2月医疗保险明细【白城分公司】", "", "", "", "", "", ""],
            ["序号", "单位名称", "已参保", "", "", "", ""],
            ["", "", "参保人数", "个人2%", "大病险150", "单位6.5%", "缴费合计"],
        ]
        fp = _InsuranceWorkbook.make(tmp_path, rows)
        infos = list_insurance_sheets(fp)
        assert len(infos) == 1
        assert infos[0].has_valid_header is True
        resolved = _resolve_insurance_columns(rows)
        assert resolved is not None
        assert resolved[0] == 1
        assert resolved[1] == 2
        cols = self._final_names(resolved, rows)
        assert cols == [
            "个人2%", "大病险150", "单位6.5%", "缴费合计",
        ]

    def test_medical_2026_9_header_true(self, tmp_path):
        rows = [
            ["2026年10月医疗保险明细【白城分公司】", "", "", "", "", "", "", "", ""],
            ["序号", "", "单位名称", "", "已参保", "", "", "", ""],
            ["", "", "", "", "参保人数", "个人2%", "大病险150", "单位6.5%", "缴费合计"],
        ]
        fp = _InsuranceWorkbook.make(tmp_path, rows)
        infos = list_insurance_sheets(fp)
        assert infos[0].has_valid_header is True
        resolved = _resolve_insurance_columns(rows)
        assert resolved is not None
        assert resolved[0] == 2
        assert resolved[1] == 4
        assert len(self._final_names(resolved, rows)) == 4

    def test_injury_base_header_true(self, tmp_path):
        rows = [
            ["2026年1月工伤保险明细【白城分公司】", "", "", ""],
            ["序号", "单位名称", "已参保", "", ""],
            ["", "", "参保人数", "单位0.16%", "缴费合计"],
        ]
        resolved = _resolve_insurance_columns(rows)
        assert resolved is not None
        assert len(self._final_names(resolved, rows)) == 2

    def test_public_fund_header_true(self, tmp_path):
        rows = [
            ["白城分公司2026年1月公积金转款明细", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["序号", "单位名称", "已参保", "", "", "", "", "", "补收", "", "", "", "转款金额", "", "", "", ""],
            ["", "", "人数", "单位12%", "个人8%", "单位5%", "个人5%", "合计", "人数", "单位12%", "个人8%", "合计", "单位12%", "个人8%", "单位5%", "个人5%", "合计", "", ""],
        ]
        resolved = _resolve_insurance_columns(rows, [(2, 7, "已参保"), (8, 11, "补收"), (12, 16, "转款金额")])
        assert resolved is not None
        cols = self._final_names(resolved, rows)
        assert len(cols) == 13
        names = set(cols)
        assert names == {
            "已参保-单位12%", "已参保-个人8%", "已参保-单位5%", "已参保-个人5%", "已参保-合计",
            "补收-单位12%", "补收-个人8%", "补收-合计",
            "转款金额-单位12%", "转款金额-个人8%", "转款金额-单位5%", "转款金额-个人5%", "转款金额-合计",
        }

    def test_social_security_base_header_true(self, tmp_path):
        rows = [
            ["白城分公司2026年1月社保转款明细", "", "", "", "", "", "", "", ""],
            ["", "", "", "", "", "", "", "", ""],
            ["序号", "单位名称", "已参保", "", "", "", "转款金额", "", "", ""],
            ["", "", "人数", "养老小计（单位+个人）", "失业小计（单位+个人）", "合计", "养老小计", "失业小计", "合计", "", ""],
        ]
        resolved = _resolve_insurance_columns(rows, [(2, 5, "已参保"), (6, 8, "转款金额")])
        assert resolved is not None
        cols = self._final_names(resolved, rows)
        assert len(cols) == 6
        names = set(cols)
        assert names == {
            "养老小计（单位+个人）", "失业小计（单位+个人）", "已参保-合计",
            "养老小计", "失业小计", "转款金额-合计",
        }

    def test_social_security_9月_header_true(self, tmp_path):
        rows = [
            ["白城分公司2026年9月社保转款明细", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["序号", "单位名称", "已参保", "", "", "", "基数调整补缴7-8月", "", "", "", "转款金额", "", "", "", "", "", "", "", ""],
            ["", "", "人数", "养老小计（单位+个人）", "失业小计（单位+个人）", "合计", "养老小计（单位+个人）", "失业小计（单位+个人）", "合计", "养老小计", "失业小计", "合计", "", "", "", "", "", "", ""],
        ]
        resolved = _resolve_insurance_columns(rows, [(2, 5, "已参保"), (6, 8, "基数调整补缴7-8月"), (9, 11, "转款金额")])
        assert resolved is not None
        cols = self._final_names(resolved, rows)
        assert len(cols) == 9
        names = set(cols)
        assert names == {
            "已参保-养老小计（单位+个人）", "已参保-失业小计（单位+个人）", "已参保-合计",
            "基数调整补缴7-8月-养老小计（单位+个人）", "基数调整补缴7-8月-失业小计（单位+个人）", "基数调整补缴7-8月-合计",
            "养老小计", "失业小计", "转款金额-合计",
        }

    def test_personnel_header_false(self, tmp_path):
        rows = [
            ["序号", "姓名", "性别", "个人编号", "身份证号", "手机号", "所属单位", "新增月份", "合同日期", "人数", "工资"],
            ["1", "张三", "", "", "220881198511056620", "13353238166", "白城市洮北区公路管理段", "202501", "202501", "1", "3000"],
        ]
        assert _is_insurance_monthly_header(rows) is False
        assert _resolve_insurance_columns(rows) is None

    def test_personnel_header_false_明细9(self, tmp_path):
        rows = [
            ["序号", "姓名", "性别", "个人编号", "身份证号", "手机号", "所属单位", "新增月份", "合同日期", "人数", "工资"],
        ]
        assert _is_insurance_monthly_header(rows) is False

    def test_header_outside_40_rows_false(self, tmp_path):
        rows = [[""] * 5 for _ in range(40)]
        rows.append(["序号", "单位名称", "已参保", "", ""])
        rows.append(["", "", "参保人数", "单位0.16%", "缴费合计"])
        assert _is_insurance_monthly_header(rows[:40]) is False

    def test_junk_sheet_false(self, tmp_path):
        rows = [
            ["无关标题", "", ""],
            ["姓名", "手机号", "身份证号"],
            ["张三", "13353238166", "220881198511056620"],
        ]
        assert _is_insurance_monthly_header(rows) is False
        assert names == {
            "已参保-养老小计（单位+个人）", "已参保-失业小计（单位+个人）", "已参保-合计",
            "基数调整补缴7-8月-养老小计（单位+个人）", "基数调整补缴7-8月-失业小计（单位+个人）", "基数调整补缴7-8月-合计",
            "养老小计", "失业小计", "转款金额-合计",
        }

    def test_personnel_header_false(self, tmp_path):
        rows = [
            ["序号", "姓名", "性别", "个人编号", "身份证号", "手机号", "所属单位", "新增月份", "合同日期", "人数", "工资"],
            ["1", "张三", "", "", "220881198511056620", "13353238166", "白城市洮北区公路管理段", "202501", "202501", "1", "3000"],
        ]
        assert _is_insurance_monthly_header(rows) is False
        assert _resolve_insurance_columns(rows) is None

    def test_personnel_header_false_明细9(self, tmp_path):
        rows = [
            ["序号", "姓名", "性别", "个人编号", "身份证号", "手机号", "所属单位", "新增月份", "合同日期", "人数", "工资"],
        ]
        assert _is_insurance_monthly_header(rows) is False

    def test_header_outside_40_rows_false(self, tmp_path):
        rows = [[""] * 5 for _ in range(40)]
        rows.append(["序号", "单位名称", "已参保", "", ""])
        rows.append(["", "", "参保人数", "单位0.16%", "缴费合计"])
        assert _is_insurance_monthly_header(rows[:40]) is False

    def test_junk_sheet_false(self, tmp_path):
        rows = [
            ["无关标题", "", ""],
            ["姓名", "手机号", "身份证号"],
            ["张三", "13353238166", "220881198511056620"],
        ]
        assert _is_insurance_monthly_header(rows) is False


class TestParseInsuranceExcel:
    def _parse(self, tmp_path: Path, rows: list[list[object]]):
        fp = _InsuranceWorkbook.make(tmp_path, rows)
        return parse_insurance_excel(fp)

    def test_medical_base_empty_amount_skipped(self, tmp_path):
        rows = [
            ["2026年2月医疗保险明细【白城分公司】", "", "", "", "", "", ""],
            ["序号", "单位名称", "已参保", "", "", "", ""],
            ["", "", "参保人数", "个人2%", "大病险150", "单位6.5%", "缴费合计"],
            ["1", "白城市洮北区公路管理段", "5.0", "439.3", "", "1427.8", "1867.1"],
            ["合计", "", "5.0", "439.3", "", "1427.8", "1867.1"],
            ["总经理审批：", "财务审核：", "制表人：张艳艳", "", "", "", ""],
        ]
        sheet = self._parse(tmp_path, rows)
        assert sheet.title == "2026年2月医疗保险明细【白城分公司】"
        assert len(sheet.rows) == 1
        assert sheet.rows[0].unit_name == "白城市洮北区公路管理段"
        assert sheet.rows[0].headcount == 5
        assert len(sheet.rows[0].fees) == 3
        assert [f.name for f in sheet.rows[0].fees] == ["个人2%", "单位6.5%", "缴费合计"]
        assert [f.amount for f in sheet.rows[0].fees] == ["439.30", "1427.80", "1867.10"]
        assert sheet.maker == "张艳艳"

    def test_public_fund_empty_columns_skipped(self, tmp_path):
        rows = [
            ["白城分公司2026年1月公积金转款明细", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["序号", "单位名称", "已参保", "", "", "", "", "", "补收", "", "", "", "转款金额", "", "", "", ""],
            ["", "", "人数", "单位12%", "个人8%", "单位5%", "个人5%", "合计", "人数", "单位12%", "个人8%", "合计", "单位12%", "个人8%", "单位5%", "个人5%", "合计", "", ""],
            ["1", "白城市洮北区公路管理段", "5", "1184", "1776", "", "", "2960", "", "", "", "", "1184", "1776", "", "", "2960", "", ""],
            ["合计", "", "5", "1184", "1776", "0", "0", "2960", "0", "0", "0", "0", "1184", "1776", "0", "0", "2960", "", ""],
            ["总经理审批：", "财务审核：", "制表：  张艳艳", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ]
        sheet = self._parse(tmp_path, rows)
        assert len(sheet.rows) == 1
        assert sheet.rows[0].headcount == 5
        assert len(sheet.rows[0].fees) == 6
        assert [f.name for f in sheet.rows[0].fees] == [
            "已参保-单位12%", "已参保-个人8%", "已参保-合计",
            "转款金额-单位12%", "转款金额-个人8%", "转款金额-合计",
        ]
        assert [f.amount for f in sheet.rows[0].fees] == [
            "1184.00", "1776.00", "2960.00", "1184.00", "1776.00", "2960.00",
        ]
        assert sheet.maker == "张艳艳"

    def test_social_security_base(self, tmp_path):
        rows = [
            ["白城分公司2026年1月社保转款明细", "", "", "", "", "", "", "", ""],
            ["", "", "", "", "", "", "", "", ""],
            ["序号", "单位名称", "已参保", "", "", "", "转款金额", "", "", ""],
            ["", "", "人数", "养老小计（单位+个人）", "失业小计（单位+个人）", "合计", "养老小计", "失业小计", "合计", "", ""],
            ["1", "白城市洮北区公路管理段", "5", "5271.85", "219.65", "5491.5", "5271.85", "219.65", "5491.5", "", ""],
            ["合计", "", "5", "5271.85", "219.65", "5491.5", "5271.85", "219.65", "5491.5", "", ""],
            ["总经理审批：", "财务审核：", "制表：  张艳艳", "", "", "", "", "", "", ""],
        ]
        sheet = self._parse(tmp_path, rows)
        assert len(sheet.rows[0].fees) == 6
        assert [f.name for f in sheet.rows[0].fees] == [
            "养老小计（单位+个人）", "失业小计（单位+个人）", "已参保-合计",
            "养老小计", "失业小计", "转款金额-合计",
        ]
        assert [f.amount for f in sheet.rows[0].fees] == [
            "5271.85", "219.65", "5491.50", "5271.85", "219.65", "5491.50",
        ]

    def test_social_security_9月(self, tmp_path):
        rows = [
            ["白城分公司2026年9月社保转款明细", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["序号", "单位名称", "已参保", "", "", "", "基数调整补缴7-8月", "", "", "", "转款金额", "", "", "", "", "", "", "", ""],
            ["", "", "人数", "养老小计（单位+个人）", "失业小计（单位+个人）", "合计", "养老小计（单位+个人）", "失业小计（单位+个人）", "合计", "养老小计", "失业小计", "合计", "", "", "", "", "", "", ""],
            ["1", "白城市洮北区公路管理段", "5", "5386.65", "224.45", "5611.1", "229.6", "9.6", "239.2", "5616.25", "234.05", "5850.3", "", "", "", "", "", "", ""],
            ["合计", "", "5", "5386.65", "224.45", "5611.1", "229.6", "9.6", "239.2", "5616.25", "234.05", "5850.3", "", "", "", "", "", "", ""],
            ["总经理审批：", "分管领导审核：", "财务审核：", "业务审核：", "制表人：张艳艳", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ]
        sheet = self._parse(tmp_path, rows)
        assert len(sheet.rows[0].fees) == 9
        assert sheet.rows[0].headcount == 5
        assert sheet.maker == "张艳艳"

    def test_more_than_17_fees_raises(self, tmp_path):
        header = ["序号", "单位名称", "已参保"] + [f"费用{i}" for i in range(18)]
        data = ["1", "单位", "5"] + ["1"] * 18
        total = ["合计", "", "5"] + ["1"] * 18
        rows = [
            ["标题", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
            ["序号", "单位名称", "已参保"] + [""] * 18,
            ["", "", "人数"] + header[3:],
            data,
            total,
            ["制表人：张三", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", ""],
        ]
        with pytest.raises(ValueError, match="费用项超过17项"):
            self._parse(tmp_path, rows)

    def test_personnel_sheet_not_valid(self, tmp_path):
        rows = [
            ["序号", "姓名", "性别", "个人编号", "身份证号", "手机号", "所属单位", "新增月份", "合同日期", "人数", "工资"],
            ["1", "张三", "", "", "220881198511056620", "13353238166", "白城市洮北区公路管理段", "202501", "202501", "1", "3000"],
        ]
        fp = _InsuranceWorkbook.make(tmp_path, rows)
        infos = list_insurance_sheets(fp)
        assert infos[0].has_valid_header is False
        with pytest.raises(ValueError, match="未找到有效分单位表"):
            parse_insurance_excel(fp)

    def test_no_data_rows_raises(self, tmp_path):
        rows = [
            ["2026年2月医疗保险明细【白城分公司】", "", "", "", "", "", ""],
            ["序号", "单位名称", "已参保", "", "", "", ""],
            ["", "", "参保人数", "个人2%", "大病险150", "单位6.5%", "缴费合计"],
            ["合计", "", "5.0", "439.3", "", "1427.8", "1867.1"],
        ]
        with pytest.raises(ValueError, match="未找到数据行"):
            self._parse(tmp_path, rows)